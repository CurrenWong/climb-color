"""把 paper/main.md 转成 DOCX（便于阅读/分享）。

支持: 标题层级 / 表格 / 代码块 / 列表 / 引用 / 粗体 / 行内代码 /
      论文图引用（检测 `paper/figN_*.png|jpg` 并嵌入图片）。
用法: python make_docx.py [输出文件名]
"""
import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = Path(__file__).parent
SRC = BASE / "main.md"
OUT = BASE / (sys.argv[1] if len(sys.argv) > 1 else "climbing_route_color_paper.docx")

CJK = '微软雅黑'
LATIN = 'Calibri'
MONO = 'Consolas'
CODE_BG = RGBColor(0x33, 0x33, 0x33)

IMG_RE = re.compile(r'paper[/\\](fig[\w]+_[\w]+\.(?:png|jpg))')


def set_style_fonts(doc):
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3',
                 'Heading 4', 'Title', 'List Bullet', 'List Number'):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = LATIN
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = rpr.makeelement(qn('w:rFonts'), {})
            rpr.append(rf)
        rf.set(qn('w:eastAsia'), CJK)
    doc.styles['Normal'].font.size = Pt(10.5)


def add_runs(par, text):
    """处理 **bold**、`code`、*italic*。"""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9.5)
            r.font.color.rgb = CODE_BG
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def parse_table(lines, i):
    """从 lines[i] 开始解析 markdown 表格，返回 (rows, next_i)。"""
    rows = []
    while i < len(lines) and lines[i].strip().startswith('|'):
        raw = lines[i].strip().strip('|')
        cells = [c.strip() for c in raw.split('|')]
        if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
            rows.append(cells)
        i += 1
    return rows, i


def main():
    lines = SRC.read_text(encoding='utf-8').split('\n')
    doc = docx.Document()
    set_style_fonts(doc)

    i = 0
    in_code = False
    code_buf = []
    n_tables = 0
    n_images = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---- code block ----
        if stripped.startswith('```'):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run('\n'.join(code_buf))
                r.font.name = MONO
                r.font.size = Pt(8.5)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # ---- blank ----
        if not stripped:
            i += 1
            continue

        # ---- horizontal rule ----
        if re.fullmatch(r'-{3,}', stripped):
            doc.add_paragraph().add_run('—' * 30).font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue

        # ---- headings ----
        m = re.match(r'^(#{1,5})\s+(.*)$', stripped)
        if m:
            lvl = len(m.group(1))
            text = m.group(2)
            if lvl == 1:
                h = doc.add_heading(level=0)
            else:
                h = doc.add_heading(level=min(lvl - 1, 4))
            add_runs(h, text)
            i += 1
            continue

        # ---- table ----
        if stripped.startswith('|'):
            rows, i = parse_table(lines, i)
            if rows:
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=ncols)
                t.style = 'Light Grid Accent 1'
                for ri, row in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(ncols):
                        txt = row[ci] if ci < len(row) else ''
                        cp = cells[ci].paragraphs[0]
                        add_runs(cp, txt)
                        for r in cp.runs:
                            r.font.size = Pt(9)
                            if ri == 0:
                                r.bold = True
                n_tables += 1
                doc.add_paragraph()
            continue

        # ---- blockquote ----
        if stripped.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            add_runs(p, stripped[2:])
            i += 1
            continue

        # ---- list ----
        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            i += 1
            continue

        # ---- paragraph (may embed a figure) ----
        p = doc.add_paragraph()
        add_runs(p, stripped)
        imgs = IMG_RE.findall(stripped)
        for img in imgs:
            ip = BASE.parent / 'paper' / img
            if ip.exists():
                doc.add_picture(str(ip), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                n_images += 1
        i += 1

    doc.save(OUT)
    print(f"saved {OUT}")
    print(f"  tables={n_tables}  images={n_images}  paragraphs={len(doc.paragraphs)}")


if __name__ == '__main__':
    main()
